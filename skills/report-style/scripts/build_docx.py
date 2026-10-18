#!/usr/bin/env python3
"""Create an editable DOCX report from the same Markdown used for PDF reports."""

# /// script
# requires-python = ">=3.10"
# dependencies = [
#   "python-docx>=1.2,<2",
#   "markdown>=3.11,<4",
#   "beautifulsoup4>=4.15,<5",
#   "playwright>=1.63,<2",
# ]
# ///

import argparse
import tempfile
import uuid
from pathlib import Path

from bs4 import BeautifulSoup, NavigableString
from build_doc import ASSETS, build_html, find_chrome
from docx import Document
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.opc.packuri import PackURI
from docx.opc.part import Part
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor
from lxml import etree

GREEN = "186752"
INK = "070707"
MUTED = "5B5F58"
TINT = "F1F2ED"
YELLOW = "EFCC3A"
FONT = "Space Grotesk"


def embed_fonts(document):
    """Keep the designated typeface available when Word opens the report elsewhere."""
    table = document.part.part_related_by(RT.FONT_TABLE)
    root = parse_xml(table.blob)
    font = next((item for item in root if item.get(qn("w:name")) == FONT), None)
    if font is None:
        font = element(root, "font", name=FONT)
    for tag, face in (("embedRegular", "Regular"), ("embedBold", "Bold")):
        key = uuid.uuid4()
        payload = bytearray((Path(ASSETS) / "fonts" / f"SpaceGrotesk-{face}.ttf").read_bytes())
        # OOXML font embedding reverses the GUID bytes and XORs the first 32 font bytes.
        mask = key.bytes[::-1]
        for index in range(32):
            payload[index] ^= mask[index % 16]
        part = Part(
            PackURI(f"/word/fonts/SpaceGrotesk-{face}.odttf"),
            "application/vnd.openxmlformats-officedocument.obfuscatedFont",
            bytes(payload),
            document.part.package,
        )
        reference = element(font, tag, fontKey="{" + str(key).upper() + "}", subsetted="false")
        reference.set(qn("r:id"), table.relate_to(part, RT.FONT))
    table._blob = etree.tostring(root, encoding="UTF-8", xml_declaration=True, standalone=True)
    element(document.settings.element, "embedTrueTypeFonts")


def element(parent, tag, **attributes):
    child = OxmlElement("w:" + tag)
    for name, value in attributes.items():
        child.set(qn("w:" + name), str(value))
    parent.append(child)
    return child


def shading(parent, color):
    element(parent, "shd", val="clear", fill=color)


def inline(paragraph, node, bold=None, italic=None, code=False):
    if isinstance(node, NavigableString):
        run = paragraph.add_run(str(node))
        run.bold, run.italic = bold, italic
        run.font.name = FONT
        if code:
            run.font.name = "DejaVu Sans Mono"
            run.font.size = Pt(7.6)
            shading(run._r.get_or_add_rPr(), TINT)
        return
    if node.name == "br":
        paragraph.add_run().add_break()
        return
    if node.name in ("ul", "ol"):
        return
    for child in node.children:
        inline(
            paragraph,
            child,
            True if node.name in ("b", "strong") else bold,
            True if node.name in ("i", "em") else italic,
            code or node.name == "code",
        )
    if node.name == "a" and node.get("href"):
        # Retain the source URL even when a reader cannot activate Word hyperlinks.
        paragraph.add_run(" (" + node["href"] + ")")


def paragraph_border(paragraph, edge, color, width=6, kind="dashed"):
    borders = element(paragraph._p.get_or_add_pPr(), "pBdr")
    element(borders, edge, val=kind, sz=width, color=color, space=8)


def configure(document):
    section = document.sections[0]
    section.page_width, section.page_height = Mm(210), Mm(297)
    section.top_margin = section.bottom_margin = Mm(17)
    section.left_margin = section.right_margin = Mm(19)
    for name in ("Normal", "Title", "Heading 1", "Heading 2", "List Bullet", "List Number"):
        style = document.styles[name]
        style.font.name = FONT
        style.font.color.rgb = RGBColor.from_string(INK)
        fonts = style._element.rPr.rFonts
        for attribute in ("asciiTheme", "hAnsiTheme", "eastAsiaTheme", "cstheme", "csTheme"):
            fonts.attrib.pop(qn("w:" + attribute), None)
        fonts.set(qn("w:eastAsia"), FONT)
        fonts.set(qn("w:cs"), FONT)
        properties = style._element.pPr
        if properties is not None:
            for border in list(properties.findall(qn("w:pBdr"))):
                properties.remove(border)
    normal = document.styles["Normal"]
    normal.font.size = Pt(9.4)
    normal.paragraph_format.line_spacing = 1.55
    normal.paragraph_format.space_after = Pt(6)
    for name, size, color, before in (
        ("Title", 21, INK, 4),
        ("Heading 1", 12.5, GREEN, 18),
        ("Heading 2", 10.2, INK, 13.5),
    ):
        style = document.styles[name]
        style.font.size, style.font.bold = Pt(size), True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(6)
        style.paragraph_format.keep_with_next = True


def cover(document, node):
    eyebrow = node.select_one(".eyebrow")
    if eyebrow:
        paragraph = document.add_paragraph()
        run = paragraph.add_run(eyebrow.get_text().upper())
        run.font.size, run.font.bold = Pt(7.6), True
        run.font.color.rgb = RGBColor.from_string(GREEN)
        paragraph.paragraph_format.keep_with_next = True
    document.add_paragraph(node.h1.get_text(), style="Title")
    paragraph = document.add_paragraph()
    meta = node.select_one(".meta")
    if meta:
        for index, span in enumerate(meta.find_all("span", recursive=False)):
            if index:
                paragraph.add_run("    ")
            for child in span.children:
                inline(paragraph, child)
            if "badge" in span.get("class", []):
                shading(paragraph.runs[-1]._r.get_or_add_rPr(), YELLOW)
    for run in paragraph.runs:
        run.font.size = Pt(8.4)
    paragraph_border(paragraph, "bottom", INK, width=8)
    paragraph.paragraph_format.keep_with_next = True


def heading(document, node):
    paragraph = document.add_paragraph(style="Heading 1" if node.name == "h2" else "Heading 2")
    number = node.select_one(".n")
    if number:
        run = paragraph.add_run(number.get_text() + "  ")
        run.font.color.rgb, run.bold = RGBColor.from_string(MUTED), False
        number.extract()
    for child in node.children:
        inline(paragraph, child)


def table(document, node):
    rows = node.find_all("tr")
    columns = max(len(row.find_all(["th", "td"], recursive=False)) for row in rows)
    result = document.add_table(rows=0, cols=columns)
    result.autofit = False
    width = Mm(172)
    weights = [
        max(
            12,
            min(
                60,
                max(
                    (
                        len(row.find_all(["th", "td"], recursive=False)[i].get_text())
                        for row in rows
                        if len(row.find_all(["th", "td"], recursive=False)) > i
                    ),
                    default=12,
                ),
            ),
        )
        for i in range(columns)
    ]
    for column, weight in zip(result.columns, weights, strict=True):
        column.width = int(width * weight / sum(weights))
    for source in rows:
        row = result.add_row()
        element(row._tr.get_or_add_trPr(), "cantSplit")
        is_header = bool(source.find("th"))
        if is_header:
            element(row._tr.get_or_add_trPr(), "tblHeader")
        for index, (cell, source_cell) in enumerate(
            zip(row.cells, source.find_all(["th", "td"], recursive=False), strict=False)
        ):
            cell.width = int(width * weights[index] / sum(weights))
            properties = cell._tc.get_or_add_tcPr()
            margins = element(properties, "tcMar")
            for edge, value in (("top", 105), ("bottom", 105), ("left", 135), ("right", 135)):
                element(margins, edge, w=value, type="dxa")
            borders = element(properties, "tcBorders")
            for edge in ("top", "left", "right"):
                element(borders, edge, val="nil")
            element(borders, "bottom", val="dashed", sz=6, color=INK if is_header else "B9BBB4")
            if is_header:
                shading(properties, TINT)
            paragraph = cell.paragraphs[0]
            for child in source_cell.children:
                inline(paragraph, child, bold=is_header or index == 0)
            paragraph.paragraph_format.line_spacing = 1.45
            paragraph.paragraph_format.space_after = Pt(0)
            for run in paragraph.runs:
                run.font.size = Pt(8.2)
    document.add_paragraph().paragraph_format.space_after = Pt(3)


def lists(document, node, depth=0):
    for item in node.find_all("li", recursive=False):
        paragraph = document.add_paragraph(
            style="List Number" if node.name == "ol" else "List Bullet"
        )
        paragraph.paragraph_format.left_indent = Mm(5 + depth * 5)
        for child in item.children:
            inline(paragraph, child)
        for nested in item.find_all(["ul", "ol"], recursive=False):
            lists(document, nested, depth + 1)


def build(source, output):
    page = build_html(str(source))
    soup = BeautifulSoup(page, "html.parser")
    document = Document()
    configure(document)
    embed_fonts(document)
    figures = soup.select(".fig")
    with tempfile.TemporaryDirectory(prefix="report-docx-") as temporary:
        images = []
        if figures:
            # Prose and tables stay editable. Diagrams retain the supplied HTML design as images.
            from playwright.sync_api import sync_playwright

            with sync_playwright() as playwright:
                browser = playwright.chromium.launch(
                    executable_path=find_chrome(),
                    headless=True,
                    args=["--no-sandbox", "--disable-dev-shm-usage"],
                )
                tab = browser.new_page(
                    viewport={"width": 820, "height": 1100}, device_scale_factor=2
                )
                tab.set_content(page, wait_until="networkidle")
                tab.wait_for_function("document.body.classList.contains('drawn')")
                for index, figure in enumerate(tab.locator(".fig").all()):
                    image = Path(temporary) / f"figure-{index}.png"
                    figure.screenshot(path=str(image))
                    images.append(image)
                browser.close()
        figure_index = 0
        for node in soup.body.children:
            if isinstance(node, NavigableString):
                continue
            if node.name == "header":
                cover(document, node)
            elif node.name in ("h2", "h3"):
                heading(document, node)
            elif node.name == "table":
                table(document, node)
            elif node.name in ("ul", "ol"):
                lists(document, node)
            elif "fig" in node.get("class", []):
                document.add_picture(str(images[figure_index]), width=Mm(172))
                figure_index += 1
            elif node.name == "pre":
                paragraph = document.add_paragraph()
                run = paragraph.add_run(node.get_text())
                run.font.name, run.font.size = "DejaVu Sans Mono", Pt(7.2)
                shading(paragraph._p.get_or_add_pPr(), TINT)
                paragraph.paragraph_format.keep_together = True
            elif "callout" in node.get("class", []):
                paragraphs = node.find_all("p", recursive=False) or [node]
                for text in paragraphs:
                    paragraph = document.add_paragraph()
                    for child in text.children:
                        inline(paragraph, child)
                    shading(paragraph._p.get_or_add_pPr(), TINT)
                    warning = any(c in ("c-warning", "c-danger") for c in node.get("class", []))
                    paragraph_border(
                        paragraph, "left", YELLOW if warning else GREEN, width=24, kind="single"
                    )
                    paragraph.paragraph_format.left_indent = Mm(4)
            elif node.name == "p":
                paragraph = document.add_paragraph()
                for child in node.children:
                    inline(paragraph, child)
        document.core_properties.title = soup.title.get_text()
        document.save(output)


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("markdown", type=Path)
    parser.add_argument("--out", required=True, type=Path)
    args = parser.parse_args()
    build(args.markdown, args.out)
    print(f"Wrote {args.out}")
